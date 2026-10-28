"""Download client-listed Veeva versions, annotations, and audit trails, organize them, and create a delivery ZIP."""
from __future__ import annotations

import argparse
import csv
import errno
import os
import shutil
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import requests

import veeva_download as downloader
import zip_creator as organizer

ROOT = Path(__file__).resolve().parent
REQUIRED_COLUMNS = {"Document ID", "Document Number", "Major Version Number", "Minor Version Number"}
DOWNLOAD_ERRORS = (requests.RequestException, RuntimeError, ValueError, OSError, KeyError)


def project_path(value: str) -> Path:
    path = Path(value)
    return (path if path.is_absolute() else ROOT / path).resolve()


def require_project_path(path: Path, label: str) -> None:
    if path == ROOT or ROOT not in path.parents:
        raise ValueError(f"Refusing to use {label} outside the project directory: {path}")


def clean_folder(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    for child in path.iterdir():
        if child.is_dir() and not child.is_symlink():
            shutil.rmtree(child)
        else:
            child.unlink()


def read_version_list(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"Version-list workbook not found: {path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["All Versions"] if "All Versions" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, ())
    columns = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
    missing = REQUIRED_COLUMNS - columns.keys()
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(sorted(missing))}")

    def version_part(value: Any) -> Any:
        return int(value) if isinstance(value, float) and value.is_integer() else value

    documents = []
    seen = set()
    for row in rows:
        document_id = row[columns["Document ID"]]
        number = row[columns["Document Number"]]
        major = version_part(row[columns["Major Version Number"]])
        minor = version_part(row[columns["Minor Version Number"]])
        if any(value is None for value in (document_id, number, major, minor)):
            continue
        key = (str(document_id), str(major), str(minor))
        if key in seen:
            continue
        seen.add(key)
        documents.append({"id": document_id, "document_number__v": str(number).strip(), "major_version_number__v": major, "minor_version_number__v": minor, "name__v": row[columns["Document Name"]] if "Document Name" in columns else "", "filename__v": row[columns["File Name"]] if "File Name" in columns else ""})
    if not documents:
        raise ValueError(f"No usable document-version rows found in {path}")
    return documents


def read_previous_manifest(path: Path) -> dict[str, dict[str, str]]:
    if not path.is_file():
        return {}
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return {f"{row['document_id']}|{row['version']}": row for row in csv.DictReader(handle)}


def write_manifest(rows: list[dict[str, str]], input_folder: Path) -> Path:
    manifest = input_folder / "veeva_download_manifest.csv"
    temporary = manifest.with_name(manifest.name + ".part")
    with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=downloader.MANIFEST_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(manifest)
    return manifest


def find_existing_download(input_folder: Path, document: dict[str, Any]) -> Path | None:
    number = str(document["document_number__v"])
    version = f"{document['major_version_number__v']}.{document['minor_version_number__v']}"
    folder = input_folder / number / "Source"
    prefix = f"{number}_v{version}_".casefold()
    if not folder.is_dir():
        return None
    return next((path for path in folder.iterdir() if path.is_file() and path.name.casefold().startswith(prefix) and path.stat().st_size > 0), None)


def nonempty(path: Path) -> bool:
    return path.is_file() and path.stat().st_size > 0


def is_storage_full(error: Exception) -> bool:
    return isinstance(error, OSError) and error.errno == errno.ENOSPC


def fetch_annotation_and_audit(api_root: str, api_version: str, session_id: str, document: dict[str, Any], input_folder: Path, result: dict[str, str], audits: dict[str, tuple[str, str]]) -> None:
    annotation = downloader.annotation_path(input_folder, document)
    if nonempty(annotation):
        result["annotation_status"], result["annotation_file"] = "DOWNLOADED", str(annotation)
    else:
        try:
            result["annotation_file"] = str(downloader.download_annotations(api_root, api_version, session_id, document, input_folder))
            result["annotation_status"] = "DOWNLOADED"
        except DOWNLOAD_ERRORS as error:
            if is_storage_full(error):
                raise
            result["annotation_status"] = f"FAILED: {error}"
    document_id = str(document["id"])
    if document_id not in audits:
        audit = downloader.audit_trail_path(input_folder, document)
        if nonempty(audit):
            audits[document_id] = ("DOWNLOADED", str(audit))
        else:
            try:
                audits[document_id] = ("DOWNLOADED", str(downloader.download_audit_trail(api_root, api_version, session_id, document, input_folder)))
            except DOWNLOAD_ERRORS as error:
                if is_storage_full(error):
                    raise
                audits[document_id] = (f"FAILED: {error}", "")
    result["audit_status"], result["audit_file"] = audits[document_id]


def run(clean: bool = False, start_index: int = 1, end_index: int | None = None, package: bool = True) -> int:
    downloader.load_env_file(ROOT / ".env")
    downloader.configure_proxy()
    input_folder = project_path(downloader.required_config("VEEVA_INPUT_FOLDER"))
    output_root = project_path(downloader.required_config("VEEVA_OUTPUT_FOLDER"))
    workbook_path = project_path(downloader.required_config("VEEVA_VERSION_LIST"))
    request_id = downloader.required_config("VEEVA_REQUEST_ID")
    require_project_path(input_folder, "input folder")
    require_project_path(output_root, "output folder")
    require_project_path(workbook_path, "version-list workbook")
    if input_folder == output_root or input_folder in output_root.parents or output_root in input_folder.parents:
        raise ValueError("Input and output folders must be separate and non-overlapping")
    if not request_id or any(char in request_id for char in '<>:"/\\|?*'):
        raise ValueError("VEEVA_REQUEST_ID must be a non-empty folder-safe name")

    documents = read_version_list(workbook_path)
    end_index = len(documents) if end_index is None else min(end_index, len(documents))
    if start_index < 1 or end_index < start_index:
        raise ValueError("Invalid version index range")
    print(f"Version list: {workbook_path}")
    print(f"Prepared {len(documents)} versions; processing rows {start_index}-{end_index}.")
    if clean:
        clean_folder(input_folder)
        clean_folder(output_root)
        print("Cleaned project input/output folders.")

    previous = read_previous_manifest(input_folder / "veeva_download_manifest.csv")
    auth_url = os.environ.get("STAGE_AUTH_URL", "").strip() or downloader.required_config("AUTH_URL")
    api_version = os.environ.get("VEEVA_API_VERSION", "v26.2").strip()
    session_id, api_root = downloader.authenticate(auth_url, downloader.required_config("VEEVA_USERNAME"), downloader.required_config("VEEVA_PASSWORD"))
    downloader.get_metadata(api_root, api_version, session_id)

    results = []
    audits: dict[str, tuple[str, str]] = {}
    storage_full = False
    for index, document in enumerate(documents, start=1):
        number = str(document["document_number__v"])
        version = f"{document['major_version_number__v']}.{document['minor_version_number__v']}"
        key = f"{document['id']}|{version}"
        in_range = start_index <= index <= end_index
        result = {"document_number": number, "document_name": str(document.get("name__v") or ""), "filename__v": str(document.get("filename__v") or ""), "version": version, "document_id": str(document["id"]), "status": "NOT_ATTEMPTED", "downloaded_file": "", "error": "", "annotation_status": "NOT_ATTEMPTED", "annotation_file": "", "audit_status": "NOT_ATTEMPTED", "audit_file": ""}
        if not in_range:
            result.update({field: value for field, value in previous.get(key, {}).items() if field in result and value})
        existing = find_existing_download(input_folder, document)
        if existing:
            result.update({"status": "DOWNLOADED", "downloaded_file": str(existing), "error": ""})
        elif in_range:
            try:
                destination = downloader.download_document_version(api_root, api_version, session_id, document, input_folder, verbose=False)
                result["status"] = "DOWNLOADED"
                result["downloaded_file"] = str(destination)
            except DOWNLOAD_ERRORS as error:
                result["status"] = "FAILED"
                result["error"] = str(error)
                print(f"[{index}/{len(documents)}] FAILED {number} v{version}: {error}")
                storage_full = is_storage_full(error)
        if in_range and not storage_full:
            try:
                fetch_annotation_and_audit(api_root, api_version, session_id, document, input_folder, result, audits)
            except OSError as error:
                print(f"[{index}/{len(documents)}] FAILED annotation/audit {number} v{version}: {error}")
                storage_full = True
            if result["annotation_status"].startswith("FAILED") or result["audit_status"].startswith("FAILED"):
                print(f"[{index}/{len(documents)}] {number} v{version}: annotation={result['annotation_status']}; audit={result['audit_status']}")
        results.append(result)
        if (in_range and index % 25 == 0) or index == end_index or storage_full:
            manifest = write_manifest(results, input_folder)
            completed = sum(row["status"] == "DOWNLOADED" for row in results)
            print(f"Checkpoint {index}/{len(documents)}: {completed} downloaded; manifest: {manifest}")
        if storage_full:
            print("Stopping this batch after local storage reported no space.")
            break

    manifest = write_manifest(results, input_folder)
    downloaded = sum(row["status"] == "DOWNLOADED" for row in results)
    annotations = sum(row["annotation_status"] == "DOWNLOADED" for row in results)
    audit_docs = {row["document_id"] for row in results}
    audit_count = len({row["document_id"] for row in results if row["audit_status"] == "DOWNLOADED"})
    print(f"Downloads: {downloaded}/{len(documents)}; manifest: {manifest}")
    print(f"Annotations: {annotations}/{len(documents)} versions; audit trails: {audit_count}/{len(audit_docs)} documents")
    if not package:
        return 0 if not storage_full else 1
    if downloaded == 0:
        raise RuntimeError("No source files downloaded; refusing to package an empty delivery")

    delivery = organizer.process(input_folder, request_id, output_root, copy_files=False)
    zip_path = delivery / f"{request_id}_delivery.zip"
    with zipfile.ZipFile(zip_path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise RuntimeError(f"ZIP integrity check failed at {bad_member}")
        entry_count = len(archive.namelist())
    print(f"Delivery folder: {delivery}")
    print(f"Delivery ZIP: {zip_path}")
    print(f"ZIP integrity: OK ({entry_count} entries)")
    print(f"Download failures or not attempted: {len(documents) - downloaded}")
    return 0 if downloaded == len(documents) and not storage_full else 1


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--clean", action="store_true", help="Clear configured project input/output contents before this run")
    parser.add_argument("--start-index", type=int, default=1, help="First 1-based workbook row to process")
    parser.add_argument("--end-index", type=int, help="Last 1-based workbook row to process")
    parser.add_argument("--no-package", action="store_true", help="Download/checkpoint only; package in a later run")
    args = parser.parse_args()
    raise SystemExit(run(args.clean, args.start_index, args.end_index, not args.no_package))


if __name__ == "__main__":
    main()
